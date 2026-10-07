# -*- coding: utf-8 -*-
"""
生成《心房搜查令》电影分镜剧照与台词字幕放映PPT模板
规格：
- 16:9 电影级宽屏（13.333 x 7.5 英寸 / 1920x1080 等比画幅）
- 11 页完整放映与舞台场控结构：
  1. 封面（片头大屏放映版）
  2-9. 8 大分镜放映页（16:9 电影画幅预留槽位、顶部悬浮信息栏、底部下沉式字幕条）
  10. 片尾致谢页（“谢谢大家” · 演职人员表 · 特别鸣谢 · 黄帝内经金句）
  11. 后台全流程场控与视听调度表（9行6列精细化表格）
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_template():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # 配色方案（黑金电影级色调）
    COLOR_BG = RGBColor(11, 14, 20)          # 极黑深蓝背景底色
    COLOR_FRAME = RGBColor(22, 28, 40)       # 预留槽位深灰底色
    COLOR_GOLD = RGBColor(230, 180, 70)      # 琥珀金标
    COLOR_WHITE = RGBColor(255, 255, 255)    # 纯白
    COLOR_SUBTITLE = RGBColor(220, 225, 235) # 字幕白
    COLOR_DIM = RGBColor(140, 148, 165)      # 辅文字灰
    COLOR_BAR_BG = RGBColor(14, 18, 26)      # 悬浮条背景深色
    COLOR_BAR_BORDER = RGBColor(50, 65, 90)  # 悬浮条微光边框
    COLOR_CARD_BG = RGBColor(16, 22, 32)     # 信息卡片背景色

    scenes_data = [
        {
            "id": 1,
            "badge": "第 01 幕 · 宿舍暗潮",
            "title": "深夜孤光下的寡淡清汤面",
            "speaker": "【赵学衡】",
            "dialogue": "“古人讲淡泊明志，咱们学医的人，肠胃清爽了，脑子才转得快。我这一天三顿白面条，舒坦！”",
            "stage_cue": "[舞台提示：深夜清冷台灯孤光，大蒜就面维持极简伪装，手捧《中药学》念念有词，内心暗潮汹涌]"
        },
        {
            "id": 2,
            "badge": "第 02 幕 · 门口对峙",
            "title": "走廊白光倾泻与手持公文夹的心理搜查官",
            "speaker": "【郝朋】",
            "dialogue": "“赵学衡，我们今晚过来，是接到全寝室的联合反映，带着心房搜查的紧急任务来的。”",
            "stage_cue": "[舞台提示：房门骤启，走廊森白强光斜切入室，公文夹在手目光如炬，赵学衡筷子僵在半空惊骇回眸]"
        },
        {
            "id": 3,
            "badge": "第 03 幕 · 宿舍交锋",
            "title": "100%打卡神话与医务室频开助眠药的矛盾特写",
            "speaker": "【程澄 & 郝朋】",
            "dialogue": "程澄：“桌面打卡执行率100%。” ｜ 郝朋：“那为什么医务室的登记本上，频繁出现你开助眠药的记录？”",
            "stage_cue": "[舞台提示：全红钩日程手账与安神助眠胶囊处方并列，双手剧烈颤抖，背景幽红心律失常折线隐现]"
        },
        {
            "id": 4,
            "badge": "第 04 幕 · 暗室铁柜",
            "title": "三把黄铜大锁与身体死守柜门的赵学衡",
            "speaker": "【赵学衡 & 郝朋】",
            "dialogue": "郝朋：“心理辅导老师签发《心房关怀通行证》。” ｜ 赵学衡：“别动！这柜子不是我的！……里面什么都没有！”",
            "stage_cue": "[舞台提示：冰蓝追光笼罩三道重锁铁柜，封条严密，赵学衡恐慌扑身死守，内心防御堡垒濒临解体]"
        },
        {
            "id": 5,
            "badge": "第 05 幕 · 书籍雪崩（高潮）",
            "title": "书籍倾覆与赵学衡跪地抱头痛哭",
            "speaker": "【赵学衡】",
            "dialogue": "“郝朋……我一页都没看完……全在这儿！我怕啊！我穷怕了，也落后怕了！我快被这些纸压死了！！”",
            "stage_cue": "[舞台提示：高潮溃堤！柜门洞开成千上万试卷书籍如雪崩倾覆，单束苍白顶光下跪地抱头号啕，心魔终决堤]"
        },
        {
            "id": 6,
            "badge": "第 06 幕 · 破防与接纳",
            "title": "抗拒安慰的应激防御与平视递上的搪瓷温茶",
            "speaker": "【赵学衡 & 林凡 & 郝朋】",
            "dialogue": "赵学衡：“别碰我！” ｜ 林凡：“我不懂，你教教我……先喝口热茶。” ｜ 郝朋：“你不是贪婪，你只是太想赢。”",
            "stage_cue": "[舞台提示：暖琥珀色柔光切入，甩开安慰的刺猬应激与平视递上的温热药茶，直面脆弱，破防走向真实接纳]"
        },
        {
            "id": 7,
            "badge": "第 07 幕 · 笛声晨光",
            "title": "横吹紫竹笛（徵音清心·角音疏肝）与撕碎苛刻清单",
            "speaker": "【郝朋 & 赵学衡】",
            "dialogue": "郝朋：“接纳自己的不完美，允许自己今天只是一棵刚刚破土的幼苗。” ｜ 赵学衡：“好！去草药园！”",
            "stage_cue": "[舞台提示：晨光穿窗洒入，紫竹笛悠扬吹奏《清心引》，苛刻作息纸条被狠狠撕碎扬向晨曦，如释重负]"
        },
        {
            "id": 8,
            "badge": "第 08 幕 · 尾声谢幕",
            "title": "四人携手鞠躬致谢与《黄帝内经》经典金句大屏",
            "speaker": "【郝朋（领衔总结）】",
            "dialogue": "“大梦初醒，心门始开。正如《黄帝内经》所言——‘恬淡虚无，真气从之；精神内守，病安从来！’”",
            "stage_cue": "[舞台提示：全场金碧辉煌漫射光，四位青年演员并肩牵手优雅鞠躬，LED大屏金字璀璨，大幕从容落下]"
        }
    ]

    # ==================== 第 1 页：封面 ====================
    slide_cov = prs.slides.add_slide(blank_layout)
    bg_cov = slide_cov.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_cov.fill.solid()
    bg_cov.fill.fore_color.rgb = COLOR_BG
    bg_cov.line.color.rgb = COLOR_BG

    # 封面微光装饰边框
    frame_cov = slide_cov.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    frame_cov.fill.solid()
    frame_cov.fill.fore_color.rgb = RGBColor(14, 18, 26)
    frame_cov.line.color.rgb = COLOR_BAR_BORDER
    frame_cov.line.width = Pt(1.5)

    tb_cov = slide_cov.shapes.add_textbox(Inches(1.3), Inches(1.2), Inches(10.733), Inches(5.1))
    tf_cov = tb_cov.text_frame
    tf_cov.word_wrap = True

    p0 = tf_cov.paragraphs[0]
    p0.text = "第十四届校园心理情景剧大赛 · 舞台大屏放映系统"
    p0.font.size = Pt(14)
    p0.font.color.rgb = COLOR_GOLD
    p0.font.bold = True
    p0.font.name = "Microsoft YaHei"
    p0.space_after = Pt(20)

    p1 = tf_cov.add_paragraph()
    p1.text = "心房搜查令"
    p1.font.size = Pt(56)
    p1.font.color.rgb = COLOR_WHITE
    p1.font.bold = True
    p1.font.name = "Microsoft YaHei"
    p1.space_after = Pt(14)

    p2 = tf_cov.add_paragraph()
    p2.text = "电影分镜剧照与台词字幕放映演示文稿 · 现场放映版"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_GOLD
    p2.font.bold = True
    p2.font.name = "Microsoft YaHei"
    p2.space_after = Pt(30)

    p3 = tf_cov.add_paragraph()
    p3.text = "创作立意：当代大学生学术囤积与冒名顶替困境 ｜ 认知解离断舍离 ｜ 中医五音情志相胜"
    p3.font.size = Pt(14.5)
    p3.font.color.rgb = COLOR_SUBTITLE
    p3.font.name = "Microsoft YaHei"
    p3.space_after = Pt(14)

    p4 = tf_cov.add_paragraph()
    p4.text = "演职班级：2024级数据科学与大数据技术班  ｜  画幅规范：16:9 电影级沉浸式全屏画幅"
    p4.font.size = Pt(13)
    p4.font.color.rgb = COLOR_DIM
    p4.font.name = "Microsoft YaHei"

    # ==================== 第 2 至 9 页：8 大分镜放映页 ====================
    # 采用标准 16:9 电影级大屏设计（13.333 x 7.5 英寸）：
    # Shape 0: 全屏背景底板
    # Shape 1: 16:9 剧照占位框（注：流水线注入剧照时将直接移除此占位框，杜绝任何遮挡！）
    # Shape 2: 顶部悬浮信息栏（Badge + 幕次 + 标题 + Scene序号）
    # Shape 3: 底部下沉悬浮字幕条（角色 + 经典台词 + 舞台调度提示）

    for sc in scenes_data:
        slide = prs.slides.add_slide(blank_layout)

        # Shape 0: 全屏背景底板
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG

        # Shape 1: 16:9 全屏电影级分镜剧照预留槽位（仅在模板未注入图片时展示占位视觉）
        slot_frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        slot_frame.fill.solid()
        slot_frame.fill.fore_color.rgb = COLOR_FRAME
        slot_frame.line.color.rgb = COLOR_BAR_BORDER
        slot_frame.line.width = Pt(1.5)
        slot_tf = slot_frame.text_frame
        slot_tf.word_wrap = True
        slot_p = slot_tf.paragraphs[0]
        slot_p.text = f"[ 16:9 电影级沉浸式分镜剧照预留槽位: scene_{sc['id']:02d}.jpg ]"
        slot_p.font.size = Pt(16)
        slot_p.font.color.rgb = COLOR_DIM
        slot_p.font.name = "Microsoft YaHei"
        slot_p.alignment = PP_ALIGN.CENTER

        # Shape 2: 顶部悬浮信息栏 (Left: 0.8", Top: 0.25", Width: 11.733", Height: 0.78")
        top_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.25), Inches(11.733), Inches(0.78))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_BAR_BG
        top_bar.line.color.rgb = COLOR_BAR_BORDER
        top_bar.line.width = Pt(1.2)
        top_tf = top_bar.text_frame
        top_tf.word_wrap = True
        top_tf.margin_left = Inches(0.25)
        top_tf.margin_top = Inches(0.1)

        top_p = top_tf.paragraphs[0]
        r_badge = top_p.add_run()
        r_badge.text = sc['badge']
        r_badge.font.size = Pt(15)
        r_badge.font.bold = True
        r_badge.font.color.rgb = COLOR_GOLD
        r_badge.font.name = "Microsoft YaHei"

        r_sep = top_p.add_run()
        r_sep.text = "  ｜  "
        r_sep.font.size = Pt(15)
        r_sep.font.color.rgb = COLOR_DIM
        r_sep.font.name = "Microsoft YaHei"

        r_title = top_p.add_run()
        r_title.text = sc['title']
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_WHITE
        r_title.font.name = "Microsoft YaHei"

        top_sub = top_tf.add_paragraph()
        top_sub.text = f"《心房搜查令》电影分镜放映系统  ·  SCENE {sc['id']:02d} / 08"
        top_sub.font.size = Pt(10.5)
        top_sub.font.color.rgb = COLOR_GOLD
        top_sub.font.name = "Microsoft YaHei"

        # Shape 3: 底部下沉悬浮台词字幕条 (Left: 0.8", Top: 5.42", Width: 11.733", Height: 1.72")
        sub_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.42), Inches(11.733), Inches(1.72))
        sub_bar.fill.solid()
        sub_bar.fill.fore_color.rgb = COLOR_BAR_BG
        sub_bar.line.color.rgb = COLOR_BAR_BORDER
        sub_bar.line.width = Pt(1.5)

        sub_tf = sub_bar.text_frame
        sub_tf.word_wrap = True
        sub_tf.margin_left = Inches(0.3)
        sub_tf.margin_top = Inches(0.18)
        sub_tf.margin_right = Inches(0.3)

        spk_p = sub_tf.paragraphs[0]
        r_spk = spk_p.add_run()
        r_spk.text = sc['speaker'] + "  "
        r_spk.font.size = Pt(14.5)
        r_spk.font.bold = True
        r_spk.font.color.rgb = COLOR_GOLD
        r_spk.font.name = "Microsoft YaHei"

        r_dlg = spk_p.add_run()
        r_dlg.text = sc['dialogue']
        r_dlg.font.size = Pt(14.5)
        r_dlg.font.bold = True
        r_dlg.font.color.rgb = COLOR_WHITE
        r_dlg.font.name = "Microsoft YaHei"
        spk_p.space_after = Pt(6)

        cue_p = sub_tf.add_paragraph()
        cue_p.text = sc['stage_cue']
        cue_p.font.size = Pt(11.5)
        cue_p.font.color.rgb = COLOR_GOLD
        cue_p.font.name = "Microsoft YaHei"

    # ==================== 第 10 页：片尾“谢谢大家”致谢幻灯片 ====================
    slide_thx = prs.slides.add_slide(blank_layout)
    bg_thx = slide_thx.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_thx.fill.solid()
    bg_thx.fill.fore_color.rgb = COLOR_BG
    bg_thx.line.color.rgb = COLOR_BG

    # 顶部致谢大标题区域
    tb_thx_title = slide_thx.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(1.8))
    tf_thx_t = tb_thx_title.text_frame
    tf_thx_t.word_wrap = True

    pthx0 = tf_thx_t.paragraphs[0]
    pthx0.text = "《心房搜查令》 · 全剧终"
    pthx0.font.size = Pt(16)
    pthx0.font.bold = True
    pthx0.font.color.rgb = COLOR_GOLD
    pthx0.font.name = "Microsoft YaHei"
    pthx0.alignment = PP_ALIGN.CENTER
    pthx0.space_after = Pt(6)

    pthx1 = tf_thx_t.add_paragraph()
    pthx1.text = "谢 谢 大 家"
    pthx1.font.size = Pt(50)
    pthx1.font.bold = True
    pthx1.font.color.rgb = COLOR_WHITE
    pthx1.font.name = "Microsoft YaHei"
    pthx1.alignment = PP_ALIGN.CENTER

    # 中部金句框
    quote_bar = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.35), Inches(10.333), Inches(0.8))
    quote_bar.fill.solid()
    quote_bar.fill.fore_color.rgb = RGBColor(18, 24, 34)
    quote_bar.line.color.rgb = COLOR_GOLD
    quote_bar.line.width = Pt(1.2)
    qtf = quote_bar.text_frame
    qtf.word_wrap = True
    qp = qtf.paragraphs[0]
    qp.text = "“恬淡虚无，真气从之；精神内守，病安从来！”  ——《黄帝内经·素问》"
    qp.font.size = Pt(15)
    qp.font.bold = True
    qp.font.color.rgb = COLOR_GOLD
    qp.font.name = "Microsoft YaHei"
    qp.alignment = PP_ALIGN.CENTER

    # 下方左右两大信息卡片
    # 左卡片：演职人员表
    card_left = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(3.35), Inches(5.25), Inches(3.35))
    card_left.fill.solid()
    card_left.fill.fore_color.rgb = COLOR_CARD_BG
    card_left.line.color.rgb = COLOR_BAR_BORDER
    card_left.line.width = Pt(1.2)
    cl_tf = card_left.text_frame
    cl_tf.word_wrap = True
    cl_tf.margin_left = Inches(0.35)
    cl_tf.margin_right = Inches(0.35)
    cl_tf.margin_top = Inches(0.25)

    cl_p0 = cl_tf.paragraphs[0]
    cl_p0.text = "【 演职人员表 · CAST & CREW 】"
    cl_p0.font.size = Pt(14)
    cl_p0.font.bold = True
    cl_p0.font.color.rgb = COLOR_GOLD
    cl_p0.font.name = "Microsoft YaHei"
    cl_p0.space_after = Pt(12)

    cast_items = [
        ("编剧 / 导演", "2024级数据科学与大数据技术班 创编组"),
        ("赵学衡 (饰)", "中医药大三医学生 · 执念释怀与自省"),
        ("郝  朋 (饰)", "班级心理委员 · 心房搜查与温情引航"),
        ("程  澄 (饰)", "室友甲 · 现实客观观察者"),
        ("林  凡 (饰)", "室友乙 · 竹笛现场演奏与五音情志调和")
    ]
    for role, desc in cast_items:
        p_c = cl_tf.add_paragraph()
        r1 = p_c.add_run()
        r1.text = f"• {role}："
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_WHITE
        r1.font.name = "Microsoft YaHei"
        r2 = p_c.add_run()
        r2.text = f" {desc}"
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_SUBTITLE
        r2.font.name = "Microsoft YaHei"
        p_c.space_after = Pt(6)

    # 右卡片：特别鸣谢与致谢寄语
    card_right = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.88), Inches(3.35), Inches(5.25), Inches(3.35))
    card_right.fill.solid()
    card_right.fill.fore_color.rgb = COLOR_CARD_BG
    card_right.line.color.rgb = COLOR_BAR_BORDER
    card_right.line.width = Pt(1.2)
    cr_tf = card_right.text_frame
    cr_tf.word_wrap = True
    cr_tf.margin_left = Inches(0.35)
    cr_tf.margin_right = Inches(0.35)
    cr_tf.margin_top = Inches(0.25)

    cr_p0 = cr_tf.paragraphs[0]
    cr_p0.text = "【 特别鸣谢 · SPECIAL THANKS 】"
    cr_p0.font.size = Pt(14)
    cr_p0.font.bold = True
    cr_p0.font.color.rgb = COLOR_GOLD
    cr_p0.font.name = "Microsoft YaHei"
    cr_p0.space_after = Pt(12)

    thanks_items = [
        ("指导单位", "校学生工作部（处） / 心理健康教育与咨询中心"),
        ("学术支持", "传统中医药五音疗疾与情志相胜指导组"),
        ("剧组班级", "2024级数据科学与大数据技术班 全体同学"),
        ("技术呈现", "16:9 沉浸式电影分镜视听放映系统"),
        ("致谢寄语", "诚挚感谢各位评委老师、现场观众的悉心指导与陪伴！")
    ]
    for role, desc in thanks_items:
        p_t = cr_tf.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"• {role}："
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_WHITE
        r1.font.name = "Microsoft YaHei"
        r2 = p_t.add_run()
        r2.text = f" {desc}"
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_SUBTITLE
        r2.font.name = "Microsoft YaHei"
        p_t.space_after = Pt(6)

    # ==================== 第 11 页：后台全流程场控与视听调度表 ====================
    slide_back = prs.slides.add_slide(blank_layout)
    bg_back = slide_back.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_back.fill.solid()
    bg_back.fill.fore_color.rgb = COLOR_BG
    bg_back.line.color.rgb = COLOR_BG

    # 场控表标题
    tb_b_title = slide_back.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.733), Inches(0.8))
    tf_b = tb_b_title.text_frame
    tf_b.word_wrap = True
    pb0 = tf_b.paragraphs[0]
    pb0.text = "《心房搜查令》后台全流程场控与视听调度表"
    pb0.font.size = Pt(20)
    pb0.font.bold = True
    pb0.font.color.rgb = COLOR_WHITE
    pb0.font.name = "Microsoft YaHei"

    pb1 = tf_b.add_paragraph()
    pb1.text = "舞台剧组全流程精细化执行表 · 覆盖 8 幕分镜角色走位、道具流转、灯光预设与音效节拍点"
    pb1.font.size = Pt(11)
    pb1.font.color.rgb = COLOR_GOLD
    pb1.font.name = "Microsoft YaHei"

    # 场控表格（9行6列：1表头 + 8幕内容）
    table_shape = slide_back.shapes.add_table(9, 6, Inches(0.6), Inches(1.15), Inches(12.133), Inches(5.95))
    tbl = table_shape.table

    col_widths = [Inches(1.3), Inches(1.4), Inches(2.3), Inches(2.2), Inches(2.3), Inches(2.633)]
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = w

    headers = ["幕次 / 场景", "登场角色", "核心道具调度", "舞台灯光设计", "现场音效与配乐", "场控与调度要点"]
    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(25, 35, 52)
        cell.text_frame.margin_left = Inches(0.08)
        cell.text_frame.margin_right = Inches(0.08)
        cell.text_frame.margin_top = Inches(0.08)
        cell.text_frame.margin_bottom = Inches(0.08)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD
        p.font.name = "Microsoft YaHei"
        p.alignment = PP_ALIGN.CENTER

    table_data = [
        ("第01幕·宿舍暗潮", "赵学衡", "简易课桌、清汤白水面、生大蒜、《中药学》书、分秒必争横幅", "昏暗清冷台灯、局部顶光", "沉重机械钟表滴答声、压迫感低音铺垫", "重点把控吃面狼吞虎咽与背书念词的强装自律感"),
        ("第02幕·门口对峙", "赵学衡、郝朋、程澄", "黑色公文夹、门框暗影道具", "走廊刺眼冷白强光斜切入室，强明暗反差", "门骤开撞击声、心跳急促骤停音效", "郝朋目光如炬具审讯威慑，赵学衡筷子定格半空"),
        ("第03幕·宿舍交锋", "赵学衡、郝朋、程澄", "100%全打钩日程手账、安神助眠胶囊处方、空药板", "桌面局部特写聚光，背景隐现心电红光", "心跳重锤加速，低沉大提琴旋律切入", "红叉账本与助眠药单无情并置，层层戳穿防线"),
        ("第04幕·暗室铁柜", "赵学衡、郝朋、程澄", "双门深绿铁皮储物柜、三把黄铜大锁、封条、通行证", "刺目冰蓝垂直追光，两侧深暗心理空间", "钥匙插入铜锁刺耳扭转声、沉重铁链落地声", "赵学衡以身死死贴挡柜门，尖锐破音抗拒"),
        ("第05幕·书籍雪崩", "赵学衡、郝朋、程澄", "成百上千全新考研讲义与空白试卷、未拆封讲义", "单束苍白顶光刺射，戏剧大冲突追光", "低沉压抑弦乐极速高潮、重锤心跳轰鸣", "高潮爆发：资料海啸倾倒，赵学衡跪地抱头痛哭"),
        ("第06幕·破防与接纳", "赵学衡、郝朋、程澄、林凡", "复古绿色搪瓷温茶杯、散落试卷手帕", "温暖琥珀色漫射光渐亮，破除室内阴霾", "弦乐由激越转入舒缓长音，呼吸声减弱", "甩手抗拒防线破除，林凡平视递茶，共情接纳"),
        ("第07幕·笛声晨光", "赵学衡、郝朋、程澄、林凡", "传统紫竹笛（配红穗）、严苛日程纸条碎片", "清晨明媚阳光大面积漫射，草药园生机绿意", "现场竹笛独奏《清心引》、空灵流水古琴", "徵音清心角音疏肝，碎纸成蝶飞扬，重获释然"),
        ("第08幕·谢幕归真", "全体演员（四人并肩）", "演出服、巨幅LED背景大屏（内经金句书法）", "剧场温暖金黄色泛光全开，辉煌落幕", "温暖宏大终曲交响、全场热烈掌声混音", "四人并肩牵手优雅鞠躬，黄帝内经经典句定格")
    ]

    for ri, row in enumerate(table_data):
        row_bg = RGBColor(18, 24, 34) if ri % 2 == 0 else RGBColor(14, 18, 26)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            cell.text_frame.margin_left = Inches(0.06)
            cell.text_frame.margin_right = Inches(0.06)
            cell.text_frame.margin_top = Inches(0.06)
            cell.text_frame.margin_bottom = Inches(0.06)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.name = "Microsoft YaHei"
            if ci == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_GOLD
                p.alignment = PP_ALIGN.CENTER
            elif ci == 1:
                p.font.color.rgb = COLOR_WHITE
                p.alignment = PP_ALIGN.CENTER
            else:
                p.font.color.rgb = COLOR_SUBTITLE

    template_filename = "心房搜查令_电影分镜剧照与台词字幕放映PPT.pptx"
    prs.save(template_filename)
    print(f"模板生成成功: {template_filename}")
    print(f"总页数: {len(prs.slides)} 页")
    return template_filename

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    build_template()

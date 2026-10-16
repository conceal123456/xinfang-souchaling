# -*- coding: utf-8 -*-
"""
生成《心房搜查令》现场大屏放映版与全流程场控分镜案 PPT。

重构优化核心：
1. 彻底解耦大屏放映与后台场控：
   - 《心房搜查令_现场大屏放映版.pptx》（10 页）：
     P1 电影海报封面、P2-P9 纯场景大图（无字纯净电影画幅）、P10 剧院谢幕致谢页。
     彻底剔除第 11 页场控表，规避现场投影切出后台表格的重大舞台事故！
   - 《心房搜查令_全流程场控与分镜案.pptx》（11 页）：
     保留第 11 页 9 行 6 列完整后台全流程场控表，供导演、灯光、音效人员推演或打印。
2. 全流程注入【演讲者备注（Notes）】：
   为每一页精准注入剧情节点、切页触发时机、灯光设计、音效配乐与放映员实操指南，
   实现双屏演讲者视图下的精准专业场控。
3. P10 致谢页金句与卡片视觉抗强光升级：
   引入微光质感深色底衬面板与琥珀金饰边，杜绝剧院聚光灯冲刷字迹，字迹锐利醒目、极具专业典雅质感。
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 详尽的放映员演讲者备注（Notes）
SLIDE_NOTES = {
    1: """【幕前准备 / 封面放映】
• 剧情节点：大赛报幕、大幕未启、观众入场。
• 触发切页时机：主持人报幕完毕，舞台灯光全暗，全场静止，赵学衡在桌前坐定吃面时，切入第 2 页（第 1 幕）。
• 舞台灯光设计：全场暗光，聚光灯缓缓准备。
• 现场音效与配乐：低沉机械钟表滴答声隐约渐入。
• 放映员实操指南：保持封面稳定放映，等待舞台开场指令，切忌误触翻页。""",

    2: """【第 1 幕：宿舍暗潮 · 寡淡清汤面与伪装】
• 剧情节点：赵学衡深夜孤光吃清汤面嚼生大蒜，强装自律背诵中药方歌。
• 触发切页时机：赵学衡大喊“进！门没锁！大晚上的谁啊？”，门口急促敲门声响起、郝朋程澄推门瞬间，切入第 3 页（第 2 幕）。
• 舞台灯光设计：昏暗清冷台灯，局部顶光打在桌面，四周陷入暗影。
• 现场音效与配乐：沉重机械钟表滴答声、压迫感低音铺垫。
• 放映员实操指南：盯紧敲门动作与推门瞬间，迅速切页，配合走廊光倾泻。""",

    3: """【第 2 幕：门口对峙 · 走廊白光与心理搜查官】
• 剧情节点：郝朋夹公文夹推门而入，走廊光切入，赵学衡筷子悬停，郝朋挑明搜查任务。
• 触发切页时机：赵学衡笑称“我就是个遵纪守法、本本分分的大三穷学生啊！”，郝朋回应“你是不是本分，看看就知道了。程澄，查一查他的桌面和公开记录”时，切入第 4 页（第 3 幕）。
• 舞台灯光设计：房门骤开，走廊刺眼冷白强光斜切入室，强明暗反差剪影。
• 现场音效与配乐：门骤开撞击声、心跳急促骤停音效。
• 放映员实操指南：两人言语交锋，在郝朋下令搜查桌面时切页。""",

    4: """【第 3 幕：宿舍交锋 · 打卡神话与助眠药单】
• 剧情节点：程澄翻出100%打卡本，郝朋直击失眠多梦与助眠药单，出示通行证移步暗室铁柜。
• 触发切页时机：郝朋展开通行证“既然书桌查不出什么，那咱们就换个地方——移步去你的专属‘大宅’看一看”，演员移步储物柜时，切入第 5 页（第 4 幕）。
• 舞台灯光设计：桌面局部特写聚光，背景隐现微弱心电红光。
• 现场音效与配乐：心跳重锤加速，低沉大提琴旋律压抑切入。
• 放映员实操指南：郝朋亮出《心房关怀通行证》并移步铁柜时切页。""",

    5: """【第 4 幕：暗室铁柜 · 三把铜锁与死守柜门】
• 剧情节点：储物柜三把黄铜大锁，赵学衡以身死死挡柜，郝朋程澄合力拉开赵学衡、钥匙入锁。
• 触发切页时机：程澄大喊“起开！”，钥匙入锁发出刺耳金属声、柜门猛然被拽开瞬间，切入第 6 页（第 5 幕）。
• 舞台灯光设计：舞台骤暗，刺目冰蓝垂直追光打在储物柜上，两侧陷入深暗。
• 现场音效与配乐：钥匙插入铜锁刺耳扭转声、沉重铁链撞击声。
• 放映员实操指南：柜门拽开、开锁金属声响起的瞬间精准切页！""",

    6: """【第 5 幕：书籍雪崩 · 资料海啸与痛哭溃堤】
• 剧情节点：柜中讲义如雪崩倾倒，未拆封资料散落一地，赵学衡瘫坐痛哭崩溃，自白冒名顶替恐惧。
• 触发切页时机：赵学衡嘶喊“我快被这些纸压死了！我快被逼疯了！！”，舞台静止抽泣，林凡携紫竹笛缓步上场时，切入第 7 页（第 6 幕）。
• 舞台灯光设计：单束苍白顶光刺射倾泻的书海，戏剧大冲突追光。
• 现场音效与配乐：低沉压抑弦乐极速高潮、重锤心跳轰鸣达到顶点。
• 放映员实操指南：高潮哭诉结束，林凡上台递步时切入温情画面。""",

    7: """【第 6 幕：破防与接纳 · 甩开应激到温茶抚慰】
• 剧情节点：赵学衡应激甩手，林凡平视共情递上陈皮麦冬茶，郝朋温和剖析执念。
• 触发切页时机：郝朋说“接纳自己的不完美，允许自己今天只是一棵刚刚破土的幼苗，不需要在一天之内长成参天大树”，赵学衡准备撕碎纸条时，切入第 8 页（第 7 幕）。
• 舞台灯光设计：舞台静默片刻，温暖琥珀色漫射光渐亮，缓缓破除室内阴霾。
• 现场音效与配乐：弦乐由激越转入舒缓长音，呼吸声逐渐趋于平稳。
• 放映员实操指南：温情对话结束，纸条扬起撕碎、竹笛将起时切页。""",

    8: """【第 7 幕：笛声晨光 · 紫竹清心与释然微笑】
• 剧情节点：纸条撕碎成蝶飞扬，林凡现场吹奏紫竹笛《清心引》，程澄约定明天包子与草药园，赵学衡释怀微笑。
• 触发切页时机：赵学衡笑称“好！去草药园！……郝委员，谢谢你们。这场搜查，把我心里发霉的角落，全都晒透了”，郝朋转向观众总结时，切入第 9 页（第 8 幕）。
• 舞台灯光设计：清晨明媚阳光大面积漫射，如草药园般充满生机绿意。
• 现场音效与配乐：现场竹笛独奏《清心引》清越响起，空灵流水古琴衬底。
• 放映员实操指南：竹笛独奏段落保持画面，台词转向全场总结时切页。""",

    9: """【第 8 幕：尾声谢幕 · 内经金句与四人鞠躬】
• 剧情节点：郝朋面向全场观众从容朗诵《黄帝内经》经典名言，四位演员并肩而立。
• 触发切页时机：郝朋朗诵完“恬淡虚无，真气从之；精神内守，病安从来！”，全场掌声响起、四人鞠躬时，切入第 10 页（致谢尾页）。
• 舞台灯光设计：剧场温暖金黄色泛光全开，辉煌明亮。
• 现场音效与配乐：温暖宏大终曲交响、全场热烈掌声混音。
• 放映员实操指南：台词落地金句念完，鞠躬掌声切入致谢页！""",

    10: """【第 10 页：致谢页 · 演职人员与特别鸣谢】
• 剧情节点：全剧终谢幕，演员向观众深情鞠躬，大屏呈现演职人员表与鸣谢金句。
• 舞台灯光设计：剧场全亮，观众席泛光。
• 现场音效与配乐：谢幕背景交响乐持续播放。
• 放映员实操指南：大屏定格展示此页，直到大幕完全闭合或主持人上台，严禁再次按翻页键！""",

    11: """【第 11 页：后台全流程场控与视听调度表（仅供后台导演与技术组查阅）】
• 本页为后台剧组全流程执行表，包含 8 幕角色走位、道具流转、灯光预设、音效节拍点与场控要点。
• 供现场导演、灯光师、音效师推演、彩排或打印场控单使用。严禁在大屏向观众放映！"""
}

def get_img_path(filename):
    paths_to_check = [
        filename,
        os.path.join("剧照", filename),
        os.path.join(os.path.dirname(__file__), "..", "剧照", filename),
        os.path.join(os.path.dirname(__file__), filename)
    ]
    for p in paths_to_check:
        if os.path.exists(p):
            return os.path.abspath(p)
    return None

def build_base_slides(prs):
    blank_layout = prs.slide_layouts[6]
    
    # 标准配色
    COLOR_BG = RGBColor(11, 14, 20)          # 极黑深蓝背景底色
    COLOR_GOLD = RGBColor(245, 205, 90)      # 亮琥珀金
    COLOR_WHITE = RGBColor(255, 255, 255)    # 纯白
    COLOR_SUBTITLE = RGBColor(220, 225, 235) # 字幕柔白
    COLOR_DIM = RGBColor(140, 148, 165)      # 辅文字灰
    COLOR_CARD_BG = RGBColor(12, 16, 24)     # 质感深色底板
    COLOR_BORDER = RGBColor(60, 75, 100)     # 微光边缘色
    COLOR_GOLD_BORDER = RGBColor(190, 150, 60)# 金色边框

    # ==================== 第 1 页：封面 (cover.jpg) ====================
    slide_cov = prs.slides.add_slide(blank_layout)
    cover_img = get_img_path("cover.jpg")
    if cover_img:
        slide_cov.shapes.add_picture(cover_img, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
    else:
        bg_cov = slide_cov.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_cov.fill.solid()
        bg_cov.fill.fore_color.rgb = COLOR_BG

    # 封面电影级优雅字幕（左右角与底部留白，完全避开海报正中“心房搜查令”金字与人物）
    tb_cov_top_l = slide_cov.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(5.0), Inches(0.4))
    tf_cov_l = tb_cov_top_l.text_frame
    tf_cov_l.word_wrap = True
    p_l = tf_cov_l.paragraphs[0]
    p_l.text = "第十四届校园心理情景剧大赛 · 舞台放映"
    p_l.font.size = Pt(11.5)
    p_l.font.color.rgb = COLOR_GOLD
    p_l.font.bold = True
    p_l.font.name = "Microsoft YaHei"

    tb_cov_top_r = slide_cov.shapes.add_textbox(Inches(7.733), Inches(0.25), Inches(5.0), Inches(0.4))
    tf_cov_r = tb_cov_top_r.text_frame
    tf_cov_r.word_wrap = True
    p_r = tf_cov_r.paragraphs[0]
    p_r.text = "2024级数据科学与大数据技术班"
    p_r.font.size = Pt(11.5)
    p_r.font.color.rgb = COLOR_SUBTITLE
    p_r.font.name = "Microsoft YaHei"
    p_r.alignment = PP_ALIGN.RIGHT

    tb_cov_bot = slide_cov.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.35))
    tf_cov_bot = tb_cov_bot.text_frame
    tf_cov_bot.word_wrap = True
    p_bot = tf_cov_bot.paragraphs[0]
    p_bot.text = "电影分镜剧照与现场放映演示文稿  ｜  《心房搜查令》  ｜  16:9 全屏画幅"
    p_bot.font.size = Pt(9.5)
    p_bot.font.color.rgb = COLOR_DIM
    p_bot.font.name = "Microsoft YaHei"
    p_bot.alignment = PP_ALIGN.CENTER

    slide_cov.notes_slide.notes_text_frame.text = SLIDE_NOTES[1]

    # ==================== 第 2 至 9 页：八大分镜纯场景图 ====================
    scenes = [
        ("scene_01.jpg", "第01幕·宿舍暗潮"),
        ("scene_02.jpg", "第02幕·门口对峙"),
        ("scene_03.jpg", "第03幕·宿舍交锋"),
        ("scene_04.jpg", "第04幕·暗室铁柜"),
        ("scene_05.jpg", "第05幕·书籍雪崩"),
        ("scene_06.jpg", "第06幕·破防与接纳"),
        ("scene_07.jpg", "第07幕·笛声晨光"),
        ("scene_08.jpg", "第08幕·尾声谢幕")
    ]

    for idx, (img_file, sc_title) in enumerate(scenes, start=1):
        slide = prs.slides.add_slide(blank_layout)
        img_path = get_img_path(img_file)
        if img_path:
            slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
        else:
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = COLOR_BG
            
        slide.notes_slide.notes_text_frame.text = SLIDE_NOTES[idx + 1]

    # ==================== 第 10 页：片尾致谢页 (ending.jpg) ====================
    slide_thx = prs.slides.add_slide(blank_layout)
    ending_img = get_img_path("ending.jpg")
    if ending_img:
        slide_thx.shapes.add_picture(ending_img, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
    else:
        bg_thx = slide_thx.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_thx.fill.solid()
        bg_thx.fill.fore_color.rgb = COLOR_BG

    # 1. 顶部金句横幅：采用质感深色微光防冲刷底衬胶囊面板
    quote_panel = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(0.22), Inches(9.733), Inches(0.58))
    quote_panel.fill.solid()
    quote_panel.fill.fore_color.rgb = COLOR_BG
    quote_panel.line.color.rgb = COLOR_GOLD_BORDER
    quote_panel.line.width = Pt(1.5)
    
    qtf = quote_panel.text_frame
    qtf.word_wrap = True
    qtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp = qtf.paragraphs[0]
    qp.text = "“恬淡虚无，真气从之；精神内守，病安从来！”  ——《黄帝内经·素问》"
    qp.font.size = Pt(13.5)
    qp.font.bold = True
    qp.font.color.rgb = COLOR_GOLD
    qp.font.name = "Microsoft YaHei"
    qp.alignment = PP_ALIGN.CENTER

    # 2. 左翼卡片：演职人员表（高质感深色底衬面板，抵御舞台光晕干扰）
    card_l_bg = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.5), Inches(4.55), Inches(5.2))
    card_l_bg.fill.solid()
    card_l_bg.fill.fore_color.rgb = COLOR_CARD_BG
    card_l_bg.line.color.rgb = COLOR_BORDER
    card_l_bg.line.width = Pt(1.0)
    
    cl_tf = card_l_bg.text_frame
    cl_tf.word_wrap = True
    cl_tf.margin_left = Inches(0.18)
    cl_tf.margin_right = Inches(0.18)
    cl_tf.margin_top = Inches(0.18)

    cl_p0 = cl_tf.paragraphs[0]
    cl_p0.text = "【 演职人员表 · CAST & CREW 】"
    cl_p0.font.size = Pt(12.5)
    cl_p0.font.bold = True
    cl_p0.font.color.rgb = COLOR_GOLD
    cl_p0.font.name = "Microsoft YaHei"
    cl_p0.space_after = Pt(10)

    cast_items = [
        ("编剧 / 导演", "2024级数据科学与大数据技术班 创编组"),
        ("赵学衡 (饰)", "中医药大三医学生 · 执念释怀与自省"),
        ("郝  朋 (饰)", "班级心理委员 · 心房搜查与温情引航"),
        ("程  澄 (饰)", "室友甲 · 现实客观观察者"),
        ("林  凡 (饰)", "室友乙 · 竹笛现场演奏与五音情志调和")
    ]
    for role, desc in cast_items:
        p_c = cl_tf.add_paragraph()
        r1 = p_c.add_run()
        r1.text = f"• {role}："
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_WHITE
        r1.font.name = "Microsoft YaHei"
        r2 = p_c.add_run()
        r2.text = f" {desc}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_SUBTITLE
        r2.font.name = "Microsoft YaHei"
        p_c.space_after = Pt(7)

    # 3. 右翼卡片：特别鸣谢与致谢寄语（高质感深色底衬面板）
    card_r_bg = slide_thx.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.233), Inches(1.5), Inches(4.55), Inches(5.2))
    card_r_bg.fill.solid()
    card_r_bg.fill.fore_color.rgb = COLOR_CARD_BG
    card_r_bg.line.color.rgb = COLOR_BORDER
    card_r_bg.line.width = Pt(1.0)
    
    cr_tf = card_r_bg.text_frame
    cr_tf.word_wrap = True
    cr_tf.margin_left = Inches(0.18)
    cr_tf.margin_right = Inches(0.18)
    cr_tf.margin_top = Inches(0.18)

    cr_p0 = cr_tf.paragraphs[0]
    cr_p0.text = "【 特别鸣谢 · SPECIAL THANKS 】"
    cr_p0.font.size = Pt(12.5)
    cr_p0.font.bold = True
    cr_p0.font.color.rgb = COLOR_GOLD
    cr_p0.font.name = "Microsoft YaHei"
    cr_p0.space_after = Pt(10)

    thanks_items = [
        ("指导单位", "校学生工作部（处） / 心理健康教育与咨询中心"),
        ("学术支持", "传统中医药五音疗疾与情志相胜指导组"),
        ("剧组班级", "2024级数据科学与大数据技术班 全体同学"),
        ("技术呈现", "16:9 沉浸式电影分镜视听放映系统"),
        ("致谢寄语", "《心房搜查令》全剧终 · 谢谢大家！诚挚感谢悉心陪伴！")
    ]
    for role, desc in thanks_items:
        p_t = cr_tf.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"• {role}："
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_WHITE
        r1.font.name = "Microsoft YaHei"
        r2 = p_t.add_run()
        r2.text = f" {desc}"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_SUBTITLE
        r2.font.name = "Microsoft YaHei"
        p_t.space_after = Pt(7)

    slide_thx.notes_slide.notes_text_frame.text = SLIDE_NOTES[10]
    return prs

def add_stage_control_slide(prs):
    blank_layout = prs.slide_layouts[6]
    COLOR_BG = RGBColor(11, 14, 20)
    COLOR_GOLD = RGBColor(245, 205, 90)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_SUBTITLE = RGBColor(220, 225, 235)

    slide_back = prs.slides.add_slide(blank_layout)
    bg_back = slide_back.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_back.fill.solid()
    bg_back.fill.fore_color.rgb = COLOR_BG
    bg_back.line.color.rgb = COLOR_BG

    # 场控表标题
    tb_b_title = slide_back.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.733), Inches(0.8))
    tf_b = tb_b_title.text_frame
    tf_b.word_wrap = True
    pb0 = tf_b.paragraphs[0]
    pb0.text = "《心房搜查令》后台全流程场控与视听调度表"
    pb0.font.size = Pt(20)
    pb0.font.bold = True
    pb0.font.color.rgb = COLOR_WHITE
    pb0.font.name = "Microsoft YaHei"

    pb1 = tf_b.add_paragraph()
    pb1.text = "舞台剧组全流程精细化执行表 · 覆盖 8 幕分镜角色走位、道具流转、灯光预设与音效节拍点（仅供后台使用）"
    pb1.font.size = Pt(11)
    pb1.font.color.rgb = COLOR_GOLD
    pb1.font.name = "Microsoft YaHei"

    # 场控表格（9行6列：1表头 + 8幕内容）
    table_shape = slide_back.shapes.add_table(9, 6, Inches(0.6), Inches(1.15), Inches(12.133), Inches(5.95))
    tbl = table_shape.table

    col_widths = [Inches(1.3), Inches(1.4), Inches(2.3), Inches(2.2), Inches(2.3), Inches(2.633)]
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = w

    headers = ["幕次 / 场景", "登场角色", "核心道具调度", "舞台灯光设计", "现场音效与配乐", "场控与调度要点"]
    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(25, 35, 52)
        cell.text_frame.margin_left = Inches(0.08)
        cell.text_frame.margin_right = Inches(0.08)
        cell.text_frame.margin_top = Inches(0.08)
        cell.text_frame.margin_bottom = Inches(0.08)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD
        p.font.name = "Microsoft YaHei"
        p.alignment = PP_ALIGN.CENTER

    table_data = [
        ("第01幕·宿舍暗潮", "赵学衡", "简易课桌、清汤白水面、生大蒜、《中药学》书、分秒必争横幅", "昏暗清冷台灯、局部顶光", "沉重机械钟表滴答声、压迫感低音铺垫", "重点把控吃面狼吞虎咽与背书念词的强装自律感"),
        ("第02幕·门口对峙", "赵学衡、郝朋、程澄", "黑色公文夹、门框暗影道具", "走廊刺眼冷白强光斜切入室，强明暗反差", "门骤开撞击声、心跳急促骤停音效", "郝朋目光如炬具审讯威慑，赵学衡筷子定格半空"),
        ("第03幕·宿舍交锋", "赵学衡、郝朋、程澄", "100%全打钩日程手账、安神助眠胶囊处方、空药板", "桌面局部特写聚光，背景隐现心电红光", "心跳重锤加速，低沉大提琴旋律切入", "红叉账本与助眠药单无情并置，层层戳穿防线"),
        ("第04幕·暗室铁柜", "赵学衡、郝朋、程澄", "双门深绿铁皮储物柜、三把黄铜大锁、封条、通行证", "刺目冰蓝垂直追光，两侧深暗心理空间", "钥匙插入铜锁刺耳扭转声、沉重铁链落地声", "赵学衡以身死死贴挡柜门，尖锐破音抗拒"),
        ("第05幕·书籍雪崩", "赵学衡、郝朋、程澄", "成百上千全新考研讲义与空白试卷、未拆封讲义", "单束苍白顶光刺射，戏剧大冲突追光", "低沉压抑弦乐极速高潮、重锤心跳轰鸣", "高潮爆发：资料海啸倾倒，赵学衡跪地抱头痛哭"),
        ("第06幕·破防与接纳", "赵学衡、郝朋、程澄、林凡", "复古绿色搪瓷温茶杯、散落试卷手帕", "温暖琥珀色漫射光渐亮，破除室内阴霾", "弦乐由激越转入舒缓长音，呼吸声减弱", "甩手抗拒防线破除，林凡平视递茶，共情接纳"),
        ("第07幕·笛声晨光", "赵学衡、郝朋、程澄、林凡", "传统紫竹笛（配红穗）、严苛日程纸条碎片", "清晨明媚阳光大面积漫射，草药园生机绿意", "现场竹笛独奏《清心引》、空灵流水古琴", "徵音清心角音疏肝，碎纸成蝶飞扬，重获释然"),
        ("第08幕·谢幕归真", "全体演员（四人并肩）", "演出服、巨幅LED背景大屏（内经金句书法）", "剧场温暖金黄色泛光全开，辉煌落幕", "温暖宏大终曲交响、全场热烈掌声混音", "四人并肩牵手优雅鞠躬，黄帝内经经典句定格")
    ]

    for ri, row in enumerate(table_data):
        row_bg = RGBColor(18, 24, 34) if ri % 2 == 0 else RGBColor(14, 18, 26)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            cell.text_frame.margin_left = Inches(0.06)
            cell.text_frame.margin_right = Inches(0.06)
            cell.text_frame.margin_top = Inches(0.06)
            cell.text_frame.margin_bottom = Inches(0.06)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.name = "Microsoft YaHei"
            if ci == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_GOLD
                p.alignment = PP_ALIGN.CENTER
            elif ci == 1:
                p.font.color.rgb = COLOR_WHITE
                p.alignment = PP_ALIGN.CENTER
            else:
                p.font.color.rgb = COLOR_SUBTITLE

    slide_back.notes_slide.notes_text_frame.text = SLIDE_NOTES[11]

def build_all_presentations():
    # 1. 现场大屏放映版（10 页，无后台场控表，零失误隐患）
    prs_screen = Presentation()
    prs_screen.slide_width = Inches(13.333)
    prs_screen.slide_height = Inches(7.5)
    build_base_slides(prs_screen)
    
    screen_path = "心房搜查令_现场大屏放映版.pptx"
    prs_screen.save(screen_path)
    print(f"✓ 现场大屏放映版生成成功: {screen_path} (共 {len(prs_screen.slides)} 页)")

    # 同步更新既有放映版文件名，杜绝历史引用失误
    prs_screen.save("心房搜查令_最终现场放映版.pptx")
    prs_screen.save("心房搜查令_电影分镜剧照与台词字幕放映PPT.pptx")
    print("✓ 同步更新历史放映文件（统一为 10 页安全放映架构）")

    # 2. 全流程场控与分镜案（11 页，含后台场控表与推演执行案）
    prs_control = Presentation()
    prs_control.slide_width = Inches(13.333)
    prs_control.slide_height = Inches(7.5)
    build_base_slides(prs_control)
    add_stage_control_slide(prs_control)
    
    control_path = "心房搜查令_全流程场控与分镜案.pptx"
    prs_control.save(control_path)
    print(f"✓ 全流程场控与分镜案生成成功: {control_path} (共 {len(prs_control.slides)} 页)")

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    build_all_presentations()

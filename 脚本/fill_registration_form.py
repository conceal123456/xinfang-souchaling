# -*- coding: utf-8 -*-
"""
完善《附件2：第十四届校园心理剧大赛参赛报名表.docx》：
1. 填入规范的参赛信息（医药信息工程学院、心房搜查令、4人、原创、负责人郝朋、联系方式、备注）；
2. 附上完整的演职人员角色对应表、剧目亮点与舞台技术保障说明；
3. 严格遵循表格与文档字体规范（宋体/仿宋/黑体，字号适中，居中对齐）。
"""

import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_font(cell, text, font_name="仿宋_GB2312", size_pt=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(14)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), "Times New Roman")
    rFonts.set(qn('w:hAnsi'), "Times New Roman")
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), "Times New Roman")
    rPr.append(rFonts)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def fill_form():
    form_path = os.path.abspath('参赛材料/附件2：第十四届校园心理剧大赛参赛报名表.docx')
    doc = docx.Document(form_path)

    # 1. 完善主表格
    table = doc.tables[0]
    # 设置第一行参赛数据
    row1 = table.rows[1]
    data_row = [
        "1",
        "医药信息工程学院",
        "心房搜查令（原创）",
        "4人",
        "郝朋（心理委员）",
        "15156008899",
        "原创心理情景剧；时长12-14分钟；含16:9大屏放映与紫竹笛独奏；见下附演职表"
    ]
    
    for ci, val in enumerate(data_row):
        cell = row1.cells[ci]
        align = WD_ALIGN_PARAGRAPH.LEFT if ci == 6 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_font(cell, val, font_name="仿宋", size_pt=9.5, align=align)

    # 清除其余空白行的冗余或保留结构
    for ri in range(2, len(table.rows)):
        for cell in table.rows[ri].cells:
            cell.text = ""

    # 清除历史运行中已追加的附加表格与附录段落（确保脚本反复运行绝对幂等）
    while len(doc.tables) > 1:
        extra_tbl = doc.tables[1]._element
        extra_tbl.getparent().remove(extra_tbl)
    
    for p in list(doc.paragraphs):
        if any(kw in p.text for kw in ["演职人员角色与分工对应表", "剧目亮点与舞台呈现说明", "知识囤积障碍", "五音疗疾"]):
            p_elem = p._element
            p_elem.getparent().remove(p_elem)

    # 2. 在表格下方添加演职人员角色对应表与参赛详情
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(12)
    p_sep.paragraph_format.space_after = Pt(4)
    run_title = p_sep.add_run("【附：《心房搜查令》演职人员角色与分工对应表】")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(12)
    run_title.bold = True
    rPr = run_title._element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), "黑体")
    rPr.append(rFonts)

    # 创建演职人员子表格 (5 行 6 列)
    cast_table = doc.add_table(5, 6)
    cast_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cast_table.autofit = False

    # 表头
    cast_headers = ["序号", "角色名称", "演员姓名", "班级 / 专业", "剧情角色定位与特征", "舞台兼任工作"]
    col_widths = [Inches(0.6), Inches(1.1), Inches(1.1), Inches(2.2), Inches(3.2), Inches(1.5)]

    for ci, (h, w) in enumerate(zip(cast_headers, col_widths)):
        cell = cast_table.cell(0, ci)
        cell.width = w
        set_cell_background(cell, "F2F4F7")
        set_cell_font(cell, h, font_name="黑体", size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    cast_data = [
        ("1", "赵学衡", "赵学衡", "2024级数据科学与大数据技术班", "大三医学生，自律表象下的知识囤积与冒名顶替综合征", "主演（执念与释怀）"),
        ("2", "郝朋", "郝朋", "2024级数据科学与大数据技术班", "班级心理委员，机敏温和的“心房搜查官”与情志引航者", "主演 / 剧组负责人"),
        ("3", "程澄", "程澄", "2024级数据科学与大数据技术班", "室友甲，直率风趣的现实观察者与同伴支持力量", "主演 / 道具与场控"),
        ("4", "林凡", "林凡", "2024级数据科学与大数据技术班", "室友乙，通达沉稳的知心伙伴，引导情志疏导与接纳自我", "主演 / 竹笛现场演奏")
    ]

    for ri, row_vals in enumerate(cast_data, start=1):
        for ci, (val, w) in enumerate(zip(row_vals, col_widths)):
            cell = cast_table.cell(ri, ci)
            cell.width = w
            align = WD_ALIGN_PARAGRAPH.LEFT if ci == 4 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_font(cell, val, font_name="仿宋", size_pt=9.5, align=align)

    # 3. 补充说明段落
    notes_p = doc.add_paragraph()
    notes_p.paragraph_format.space_before = Pt(10)
    notes_p.paragraph_format.space_after = Pt(2)
    notes_p.paragraph_format.line_spacing = Pt(16)
    
    notes_text = (
        "【剧目亮点与舞台呈现说明】\n"
        "1. 剧目立意：聚焦大学校园高频心理痛点，借《人民的名义》搜查框架隐喻潜意识探索，直击“知识囤积障碍”与“冒名顶替综合征”，具有极高艺术感染力与普遍现实共鸣；\n"
        "2. 中医特色：深度融入中医药大学专业特色，引用《黄帝内经》“思则气结、土壅木郁”病机，结合“五音疗疾”（现场紫竹笛独奏《清心引》）化解气滞，展示传统医学情志调护智慧；\n"
        "3. 技术规范：表演时长严格控制在 12-14 分钟以内；配套全套 16:9 高清大屏现场放映 PPT（含双屏演讲者备注精准提示），实现灯光、音效、台词与大屏视听的高度同步。"
    )
    
    r_notes = notes_p.add_run(notes_text)
    r_notes.font.name = "Times New Roman"
    r_notes.font.size = Pt(9.5)
    r_notes.font.color.rgb = RGBColor(80, 80, 80)
    rPr_n = r_notes._element.get_or_add_rPr()
    rFonts_n = OxmlElement('w:rFonts')
    rFonts_n.set(qn('w:eastAsia'), "仿宋")
    rPr_n.append(rFonts_n)

    doc.save(form_path)
    print(f"✓ 报名表更新成功: {form_path}")

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    fill_form()

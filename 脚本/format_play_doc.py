# -*- coding: utf-8 -*-
"""
精准格式化《心房搜查令》剧本 DOCX，严格契合《附件3：第十四届校园心理剧大赛剧本格式说明(学院).docx》标准公文规范：
1. 剧本名称：方正小标宋简体二号（22pt），居中，标明（原创），行距固定值 32 磅；
2. 正文汉字：仿宋_GB2312（仿宋）三号（16pt），英文字母与数字 Times New Roman 三号（16pt），两端对齐，首行缩进 2 字符；
3. 结构层级：
   - 一级标题：黑体三号（16pt），行距固定值 32 磅，首行缩进 2 字符；
   - 二级标题：楷体_GB2312（楷体）三号（16pt），行距固定值 32 磅，首行缩进 2 字符；
   - 三级标题：仿宋_GB2312（仿宋）加粗三号（16pt），行距固定值 32 磅，首行缩进 2 字符；
4. 行距硬性要求：标题行距固定值 32 磅，正文及其他部分行距固定值 28 磅；
5. 在关键节点精准嵌入舞台场控提示：【PPT 切第 X 幕：...】、【舞台灯光：...】、【现场音效：...】；
6. 自动导出为官方标准命名 PDF：剧本/2024级数据科学与大数据技术班_心房搜查令（原创）.pdf；
7. 同步更新纯文本剧本：剧本/心房搜查令_台词文本.txt。
"""

import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client
import pythoncom

def set_run_font(run, font_ea, font_ascii="Times New Roman", size_pt=16, bold=False, color_rgb=None):
    run.font.name = font_ascii
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._element.get_or_add_rPr()
    # Remove existing rFonts if any
    for existing_rf in rPr.findall(qn('w:rFonts')):
        rPr.remove(existing_rf)
        
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_ascii)
    rFonts.set(qn('w:hAnsi'), font_ascii)
    rFonts.set(qn('w:eastAsia'), font_ea)
    rFonts.set(qn('w:cs'), font_ascii)
    rPr.append(rFonts)

def set_p_format(p, line_spacing_pt, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(line_spacing_pt)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = align
    
    if indent_chars > 0:
        p.paragraph_format.first_line_indent = Pt(16 * indent_chars)
        pPr = p._element.get_or_add_pPr()
        for old_ind in pPr.findall(qn('w:ind')):
            pPr.remove(old_ind)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        ind.set(qn('w:firstLine'), str(int(16 * indent_chars * 20)))
        pPr.append(ind)
    else:
        p.paragraph_format.first_line_indent = Pt(0)
        pPr = p._element.get_or_add_pPr()
        for old_ind in pPr.findall(qn('w:ind')):
            pPr.remove(old_ind)

def create_script_document():
    doc = docx.Document()
    
    # 页面边距：标准 A4 公文边距（上 2.54cm，下 2.54cm，左 3.17cm，右 3.17cm）
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    
    # 数据定义：(type, text)
    # type: 'title', 'h1', 'h2', 'h3', 'body'
    content = [
        # 剧本名称
        ('title', '心房搜查令（原创）'),
        
        # 一级标题
        ('h1', '一、剧目基本信息'),
        
        # 二级标题
        ('h2', '（一）创作主旨与心理学立意'),
        ('body', '本剧紧扣当代大学生的群体心理痛点，以《人民的名义》反腐搜查名场面为戏剧框架，创造性转化为一出“潜意识心房搜查”的心理情景剧。主人公赵学衡表面生活克制极简、作息严苛、云淡风轻，内心深处却饱受原生家庭匮乏感带来的“学术与资料囤积障碍”（Hoarding Disorder）以及“冒名顶替综合征”（Impostor Syndrome）折磨。'),
        ('body', '剧中通过室友兼心理委员郝朋化身“心理搜查员”层层递进的探访与搜查，彻底揭开赵学衡藏在柜子深处的虚荣打卡、泛滥网课与未拆封资料，生动再现从“强装镇定”到“心理防线总崩溃”的戏剧高潮。随后，剧本引入中医学“思则气结、土壅木郁”的病机理论与“五音疗疾”养生智慧，结合现代认知行为疗法（CBT）的认知重塑与心理着陆技术，引导学生打破病态内耗，建立接纳自我与从容求学的健康心态，兼具艺术感染力与鲜明的现实心理指导价值。'),
        
        ('h2', '（二）角色配置与演员表（共4人）'),
        ('body', '赵学衡（男，20岁）：中医药院校大三班干部、模范生。表面清心寡欲、极度规律，实则极度焦虑自卑、过度补偿式囤积学业资料。'),
        ('body', '郝朋（男，20岁）：赵学衡的同寝室友兼班级心理委员。细心机敏、温和坚定，扮演步步深入的“心房搜查官”。'),
        ('body', '程澄（男，20岁）：室友甲。性格开朗直率，代表现实生活维度的伙伴与潜意识观察者。'),
        ('body', '林凡（男，20岁）：室友乙，传统器乐特长生。沉稳通达，代表内心秩序与情志调和者（兼舞台竹笛现场演奏）。'),
        
        ('h2', '（三）舞台道具、灯光与音效配置'),
        ('h3', '1. 核心道具：'),
        ('body', '一张课桌、两把椅子、一碗清汤挂面、一套泛黄的大蒜与蒜泥碟；一个贴满封条的双门储物柜；成箱打印装订的空白试卷、厚厚的未拆封网课讲义；数十本画满红叉的打卡手账本；褪黑素空瓶与助眠冲剂盒若干；一张写满苛刻打卡目标的泛黄信纸；一只搪瓷茶杯；一根传统竹笛。'),
        ('h3', '2. 舞台灯光设计：'),
        ('body', '前半场为压抑清冷的白光与顶光，营造审讯与对峙感；高潮崩解段落切换为急速闪烁的冷色追光；心理释怀与和解阶段渐变为温馨柔和的暖黄色漫射光。'),
        ('h3', '3. 现场音效与配乐：'),
        ('body', '前期伴随沉重钟表滴答声与压迫感心跳音效；心理崩塌时配急促低沉的大提琴声；情绪疏导阶段转入现场竹笛独奏《清心引》（曲目为剧情虚构，亦可替换为《姑苏行》等传统笛曲）及流水古琴背景音。'),
        
        # 二级标题与正文
        ('h1', '二、分场分镜剧本正文'),
        
        # 第一幕
        ('h2', '（一）第一幕：宿舍暗潮——深夜孤光下的寡淡清汤面'),
        ('body', '【PPT 切第 1 页：封面海报（心房搜查令 · 舞台放映）】'),
        ('body', '【幕前准备：全场灯光暗下，聚光灯缓缓亮起，低沉钟表滴答声隐约渐入。】'),
        ('body', '【PPT 切第 2 页 / 第 1 幕：宿舍暗潮】'),
        ('body', '【舞台灯光：昏暗清冷台灯，局部顶光打在书桌，四周隐入暗影。】'),
        ('body', '【现场音效：沉重机械钟表滴答声、压迫感低音铺垫。】'),
        ('body', '【舞台布景：宿舍一角。一张简陋书桌，一盏昏暗台灯。桌上摆着一碗没有任何浇头的清汤挂面，旁边放着一瓣剥开的大蒜。墙上挂着一张写着“分秒必争”的标语。】'),
        ('body', '【人物登场：赵学衡身穿洗得发白的旧外套，大口吸溜着面条，时不时就一口大蒜。他一手端碗，一手捧着《中药学》课本，嘴里念念有词。房门突然被敲响，急促而有力。】'),
        ('body', '赵学衡（口中嚼着蒜，头也不抬大喊）：进！门没锁！大晚上的谁啊？'),
        
        # 第二幕
        ('h2', '（二）第二幕：门口对峙——走廊白光倾泻与心理搜查官'),
        ('body', '【PPT 切第 3 页 / 第 2 幕：门口对峙】'),
        ('body', '【舞台灯光：房门骤开，走廊刺眼冷白强光斜切入室，形成强烈明暗剪影对峙反差。】'),
        ('body', '【现场音效：门骤开撞击声、急促心跳骤停音效。】'),
        ('body', '【人物登场：郝朋领着程澄推门而入。郝朋夹着黑色公文夹，目光锐利如搜查官；程澄神情复杂。赵学衡筷子定格半空。】'),
        ('body', '郝朋：赵学衡同学，打扰了。这么晚还在用功呢？'),
        ('body', '赵学衡（愣了一下，随即咧嘴憨笑，站起身热络地招呼）：哎哟，郝委员！程澄！看你们这阵势，我还以为查寝的卫生督查组来了呢！来来来，坐坐坐，宿舍地方挤，没个像样的椅子。快坐！'),
        ('body', '郝朋（环视四周，冷冷一笑）：赵大学霸，你这晚饭可真是别具一格啊。清汤白水煮挂面，就着生大蒜，连点油星子都看不见。'),
        ('body', '赵学衡（端着海碗，神色自若地吸溜了一大口）：嗨！这有什么。我们老家就这习惯，大蒜杀菌通气，白面养胃安神。古人讲淡泊明志，咱们学医的人，肠胃清爽了，脑子才转得快。我这人没别的心思，一天三顿白面条，就着中药方歌，舒坦！'),
        ('body', '程澄（走上前，拿起桌角一份薄薄的作息表）：早晨5点半晨跑背方剂，晚上11点准时熄灯闭门。学衡，连辅导员都到处夸你是全院的自律标杆、作息劳模啊。'),
        ('body', '赵学衡（连连摆手，满脸堆笑）：过誉了，同志们，真的是过誉了！我算什么劳模啊。在咱们这大学里，一板砖拍下去能砸倒一大片国家奖学金获得者。我天资愚钝，唯有靠笨鸟先飞、按部就班。怎么，大晚上的突然来我这儿，不会是咱们班的学风建设要找我树典型吧？'),
        ('body', '郝朋（停下脚步，凝视赵学衡）：树典型？赵学衡，我们今晚过来，是接到全寝室的联合反映，带着心房搜查的紧急任务来的。'),
        ('body', '赵学衡（面色微变，筷子在碗边碰出清脆响声，随即故作轻松地笑起来）：搜查？搜查我？哈哈哈哈！郝委员，你大晚上的真会开玩笑！我这光溜溜的桌子，一张行军床，除了几本翻卷边的教材，我有什么好搜查的？我一不违纪二不旷课，我就是个遵纪守法、本本分分的大三穷学生啊！'),
        
        # 第三幕
        ('h2', '（三）第三幕：宿舍交锋——秩序表象下的裂纹与助眠药单'),
        ('body', '【PPT 切第 4 页 / 第 3 幕：宿舍交锋】'),
        ('body', '【舞台灯光：桌面局部特写聚光，背景隐现微弱心电红光，营造审讯心理压迫感。】'),
        ('body', '【现场音效：心跳重锤加速，低沉大提琴旋律压抑切入。】'),
        ('body', '郝朋：你是不是本分，看看就知道了。程澄，查一查他的桌面和公开记录。'),
        ('body', '【程澄依言搜查，翻开桌面整整齐齐的收纳盒，取出一本极薄的笔记本。】'),
        ('body', '程澄：报告郝委员，桌面上发现近期学习打卡记录一份。每天记录完成4门功课、阅读15篇文献、刷题50道。工工整整，全部打着红色的对钩，执行率100%。'),
        ('body', '赵学衡（昂起头，语气底气十足）：听到没有！打铁还得自身硬！班级和老师信任我，选我当干部、当榜样，我能辜负大家的信任吗？正因为我这个人原则性极强、对自己要求近乎苛刻，经常有人在背后议论我装模作样、不食人间烟火。唉，习惯了！郝朋，你们真没吃晚饭吧？这面汤还有半锅，要不给你们盛两碗垫垫肚子？'),
        ('body', '郝朋（缓步走到赵学衡面前，直击其眼神）：赵处长……哦不，赵学霸，真不用客气，这面汤你还是自己留着暖胃吧。但是，每天100%打卡完成的你，为什么这半个月来每天凌晨两点还在床上翻来覆去、唉声叹气？为什么上周的临床辩证模拟考，你答题时双手颤抖得连笔都握不住？为什么医务室的登记本上，你的名字频繁出现在开助眠药的记录里？'),
        ('body', '赵学衡（笑容瞬间僵在脸上，握着筷子的手微微颤抖，随即抬高音量）：郝朋！你这是侵犯隐私！偶尔失眠多梦，在医学生身上不正常吗？临近期中压力大，谁没个心慌气短？你们要是听信外面的风言风语，跑来我这里捕风捉影，那是对我名誉的极大侮辱！我要求你们立刻向我道歉！'),
        ('body', '郝朋（平静地掏出一张折叠整齐的纸展开）：赵学衡，这是班委会和心理辅导老师共同签发的《心房关怀通行证》。既然书桌查不出什么，那咱们就换个地方——移步去你的专属‘大宅’看一看。'),
        
        # 第四幕
        ('h2', '（四）第四幕：暗室铁柜——三把黄铜大锁与心理防线'),
        ('body', '【PPT 切第 5 页 / 第 4 幕：暗室铁柜】'),
        ('body', '【舞台灯光：舞台骤暗，刺目冰蓝垂直追光打在舞台右侧深绿铁皮储物柜上，两侧陷入深暗心理空间。】'),
        ('body', '【现场音效：钥匙插入铜锁刺耳扭转声、沉重铁链撞击声。】'),
        ('body', '【舞台布景：柜门上挂着三把黄铜大锁，贴着醒目的封条。】'),
        ('body', '赵学衡（瞳孔猛烈放大，身体下意识后退，踉跄一步靠在桌沿上）：去……去那儿干什么？那柜子放的都是些旧杂物！连我自己都快忘了里面有什么了！'),
        ('body', '郝朋（大步走到储物柜前，举起从赵学衡抽屉翻出的钥匙）：忘了？前天夜里一点，全楼道监控清清楚楚拍到你轻手轻脚打开这柜子，坐在地上抱头痛哭了半个小时。来都来了，赵大学霸，打开让大伙开开眼界吧！'),
        ('body', '赵学衡（猛扑上去死死挡在柜门前，以身护柜，声音尖锐而破音）：别动！谁也别动！这柜子不是我的！……不对，这里面什么都没有！这柜子太贵重了，我怎么可能拥有这种柜子！我想都不敢想！'),
        ('body', '程澄：你是不敢想，但你都敢做！起开！'),
        ('body', '【程澄与郝朋合力拉开赵学衡，钥匙入锁，伴随着刺耳的开锁金属声，柜门被猛然拽开！】'),
        
        # 第五幕
        ('h2', '（五）第五幕：书籍雪崩——柜门洞开与资料海啸崩塌'),
        ('body', '【PPT 切第 6 页 / 第 5 幕：书籍雪崩】'),
        ('body', '【舞台灯光：单束苍白顶光刺射倾泻的书海，戏剧大冲突追光。】'),
        ('body', '【现场音效：低沉压抑弦乐极速高潮、重锤心跳轰鸣达到顶点。】'),
        ('body', '【舞台震撼视觉呈现：成百上千份厚重的教材、考研名师讲义、竞赛题集如雪崩般从柜中倾倒而出，散落一地；紧接着掉落的是几十个未开封的打卡笔记本、十几个空了的褪黑素瓶子，以及一张写满“今日必须刷完100道题，否则就是废物”的泛黄信纸。】'),
        ('body', '郝朋（随手捡起散落在地上的讲义，撕开塑料包装）：塑封都没拆！《全国中医药综合考研真题名师点睛》——全新！《高级生物统计学精要》——全新！《名医验案全集二十讲》——全新！赵学衡，这就是你向大家展现的充实与自律？'),
        ('body', '赵学衡（双膝发软，扑通一声瘫坐在满地的书堆与纸张之中，双手抓着散落的书页，浑身剧烈颤抖，眼泪夺眶而出）：郝朋……我一页都没看完……全在这儿……我一页都没看完啊！！'),
        ('body', '郝朋（蹲下身，直视他的眼睛）：既然看不完，为什么要像疯了一样买？为什么要整箱整箱地囤？网盘里存了几个T的盗版网课，书桌下塞满了名师笔记，你甚至把一个月的生活费全花在这些从不打开的资料上，而自己天天在寝室就着大蒜啃冷馒头喝面汤？！'),
        ('body', '赵学衡（崩溃哭嚎，捶胸顿足，情绪彻底爆发）：我不敢看！但我更不敢不买啊！我怕啊！郝委员，我们家祖祖辈辈都是山沟沟里的农民，好不容易供出一个大学生，我是全村的指望！我考到这个大学，看到身边的同学，有的高中有竞赛底子，有的一口流利的专业外语，有的家里长辈就是名医……我什么都没有！我穷怕了，我也落后怕了！'),
        ('body', '赵学衡（抓起自己的头发，痛哭失声）：只要我一天不买资料、一天不打卡、有一分钟不坐在书桌前，我就觉得自己马上要被淘汰了！我就觉得自己是个冒牌货，早晚有一天会被所有人戳穿！我只有把这些书买回来、把柜子塞得满满当当，我才能骗自己——‘看啊，我也在努力，我没有落后’！可是我一打开它们，我就头疼、胸闷、喘不过气……我快被这些纸压死了！我快被逼疯了！！'),
        
        # 第六幕
        ('h2', '（六）第六幕：破防与接纳——甩开安慰到温茶抚慰'),
        ('body', '【PPT 切第 7 页 / 第 6 幕：破防与接纳】'),
        ('body', '【舞台灯光：舞台静默片刻，温暖琥珀色漫射光渐亮，缓缓破除室内阴霾。】'),
        ('body', '【现场音效：弦乐由激越转入舒缓长音，呼吸声逐渐趋于平稳。】'),
        ('body', '【人物登场：林凡携紫竹笛缓步走上舞台，伸手去扶赵学衡。】'),
        ('body', '赵学衡（猛地甩开林凡的手，声音嘶哑）：别碰我！……你们是不是觉得特别痛快？看着我这样。我不需要你们可怜我。你们根本不懂——'),
        ('body', '林凡（没有恼，慢慢蹲下身，与他平视）：我不懂。你教教我。'),
        ('body', '【长久的沉默。赵学衡的呼吸一点点平下来，肩膀不再抖了。】'),
        ('body', '林凡（这才递过一杯温热的陈皮麦冬茶）：学衡，先喝口热茶，顺顺气。'),
        ('body', '郝朋（坐在赵学衡身边的书堆旁，语气由严肃转为无比温和的共情）：学衡，你囤那些资料的时候，是把它们当成救命的稻草。可现在呢？你被它们压得喘不过气。你不是贪婪。你只是太想赢，太怕跌下来。'),
        
        # 第七幕
        ('h2', '（七）第七幕：笛声晨光——紫竹清心与认知解离断舍离'),
        ('body', '【PPT 切第 8 页 / 第 7 幕：笛声晨光】'),
        ('body', '【舞台灯光：清晨明媚阳光大面积漫射，如草药园般充满生机与暖意。】'),
        ('body', '【现场音效：现场竹笛独奏《清心引》清越响起，空灵流水古琴衬底。】'),
        ('body', '林凡（环顾满地的杂乱）：学衡，咱们在中医药大学读了三年，你还记得《素问》里怎么说的吗？‘怒伤肝，思伤脾，思则气结’。你把成山的外界焦虑一股脑塞进心里，这在医学上叫‘饮食自倍，肠胃乃伤’。精神也是有消化限度的！无节制地索取信息，非但成不了栋梁，反而导致‘宿食积滞、土壅木郁’，最后虚不受补，元气大伤啊。'),
        ('body', '赵学衡（捧着热茶，眼神茫然）：宿食积滞……虚不受补……我一直在逼自己吞下消化不了的东西……'),
        ('body', '郝朋：对！心理学上把这种行为倾向叫作‘知识囤积’——它还不是正式诊断，但和‘囤积障碍’是同一个机制。而你那种总觉得会被戳穿的感觉，叫‘冒名顶替综合征’。你总觉得现在的自己不配站在聚光灯下，必须靠堆砌外在符号来伪装坚强。但真正的强大，不是看你的网盘有多大，不是看你桌上的讲义有多厚，而是敢于直面自己当下的局限，从容地走好每一步。'),
        ('body', '郝朋（捡起那张写着苛刻打卡目标的纸条，递到赵学衡手中）：认知心理学第一步，叫作‘认知解离’与‘断舍离’。今天，当着所有人的面，把它撕了。接纳自己的不完美，允许自己今天只是一棵刚刚破土的幼苗，不需要在一天之内长成参天大树。'),
        ('body', '【赵学衡看着纸条，手颤抖着，在郝朋、程澄、林凡充满鼓励与信任的目光注视下，深吸一口气，将纸条狠狠撕成碎片，扬在空中！】'),
        ('body', '【才艺展示节点：舞台暖光普照。林凡横笛唇边，悠扬悠远的竹笛声清越响起，吹奏传统养生古曲《清心引》。笛声如清泉流淌，涤荡胸臆，营造“五音入五脏、徵音清心、角音疏肝”的宁静意境。】'),
        ('body', '程澄（走上前，拍了拍赵学衡的肩膀）：以后别一个人吃凉面生大蒜了。明天早上七点，咱们一起去食堂吃热腾腾的包子，一起去草药园慢跑，如何？'),
        ('body', '赵学衡（长长舒出一口气，整个人前所未有地松弛下来，脸上露出了久违的清澈微笑）：好！去草药园！……郝委员，谢谢你们。这场搜查，把我心里发霉的角落，全都晒透了。'),
        
        # 第八幕
        ('h2', '（八）第八幕：尾声谢幕——大屏金句定格与四人归真'),
        ('body', '【PPT 切第 9 页 / 第 8 幕：尾声谢幕】'),
        ('body', '【舞台灯光：剧场温暖金黄色泛光全开，辉煌落幕。】'),
        ('body', '【现场音效：温暖宏大的终曲交响、全场热烈掌声混音。】'),
        ('body', '郝朋（面向全场观众，从容总结）：大梦初醒，心门始开。人生的考场上，我们唯一不能敷衍的，是自己的本心；我们最需要坚守的底线，是内心的从容与坦荡。正如《黄帝内经》所言——“恬淡虚无，真气从之；精神内守，病安从来！”'),
        ('body', '【PPT 切第 10 页：全剧终致谢（演职人员表与特别鸣谢）】'),
        ('body', '【笛声渐止，四位演员在温暖明亮的舞台中央并肩而立，面带微笑，向全场观众深情鞠躬。全剧终。】')
    ]
    
    # 逐段写入格式化内容
    for item_type, text in content:
        p = doc.add_paragraph()
        run = p.add_run(text)
        
        if item_type == 'title':
            # 剧本名称：使用方正小标宋简体二号，居中，行距 32 磅
            set_run_font(run, font_ea="方正小标宋简体", size_pt=22, bold=False)
            set_p_format(p, line_spacing_pt=32, align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0)
            
        elif item_type == 'h1':
            # 一级标题：黑体三号，行距 32 磅，首行缩进 2 字符
            set_run_font(run, font_ea="黑体", size_pt=16, bold=False)
            set_p_format(p, line_spacing_pt=32, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2)
            
        elif item_type == 'h2':
            # 二级标题：楷体三号，行距 32 磅，首行缩进 2 字符
            set_run_font(run, font_ea="楷体", size_pt=16, bold=False)
            set_p_format(p, line_spacing_pt=32, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2)
            
        elif item_type == 'h3':
            # 三级标题：仿宋三号(加粗)，行距 32 磅，首行缩进 2 字符
            set_run_font(run, font_ea="仿宋", size_pt=16, bold=True)
            set_p_format(p, line_spacing_pt=32, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2)
            
        else: # body
            # 正文及其他部分：汉字仿宋，英文数字 Times New Roman，三号，行距 28 磅，首行缩进 2 字符
            set_run_font(run, font_ea="仿宋", size_pt=16, bold=False)
            set_p_format(p, line_spacing_pt=28, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2)
            
    out_docx = os.path.abspath('剧本/心房搜查令.docx')
    doc.save(out_docx)
    print(f"✓ 规范排版 DOCX 已生成: {out_docx}")
    
    # 同步输出纯文本台词（剔除行号索引，保持专业干净剧本文本）
    out_txt = os.path.abspath('剧本/心房搜查令_台词文本.txt')
    with open(out_txt, 'w', encoding='utf-8') as f:
        for item_type, text in content:
            if item_type in ('title', 'h1'):
                f.write(f"\n{'='*50}\n{text}\n{'='*50}\n\n")
            elif item_type == 'h2':
                f.write(f"\n{text}\n{'-'*35}\n")
            elif item_type == 'h3':
                f.write(f"\n{text}\n")
            else:
                f.write(f"{text}\n")
    print(f"✓ 同步纯文本台词已更新: {out_txt}")
    
    # 导出为官方提交标准 PDF
    out_pdf = os.path.abspath('剧本/2024级数据科学与大数据技术班_心房搜查令（原创）.pdf')
    export_to_pdf(out_docx, out_pdf)
    print(f"✓ 官方提交标准 PDF 已导出: {out_pdf}")

def ensure_fonts_registered():
    """确保方正小标宋等公文字体在当前 Windows 用户环境中动态加载与注册"""
    import winreg, ctypes
    font_path = os.path.expanduser('~') + r'\AppData\Local\Microsoft\Windows\Fonts\FZXBSJW.TTF'
    if os.path.exists(font_path):
        font_path = os.path.normpath(font_path)
        try:
            key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r'Software\Microsoft\Windows NT\CurrentVersion\Fonts', 0, winreg.KEY_SET_VALUE)
            winreg.SetValueEx(key, '方正小标宋简体 (TrueType)', 0, winreg.REG_SZ, font_path)
            winreg.CloseKey(key)
        except Exception:
            pass
        try:
            ctypes.windll.gdi32.AddFontResourceW(font_path)
            ctypes.windll.user32.SendMessageTimeoutW(0xFFFF, 0x001D, 0, 0, 0, 1000, 0)
        except Exception:
            pass

def export_to_pdf(docx_path, pdf_path):
    ensure_fonts_registered()
    pythoncom.CoInitialize()
    word_app = None
    try:
        word_app = win32com.client.DispatchEx("Word.Application")
        word_app.Visible = False
        doc = word_app.Documents.Open(docx_path, ReadOnly=False)
        # 开启字体嵌入与子集内嵌，确保跨机器兼容
        try:
            doc.EmbedTrueTypeFonts = True
            doc.SaveSubsetFonts = True
            doc.DoNotEmbedSystemFonts = False
            doc.Save()
        except Exception as e_embed:
            print(f"  ! 嵌入字体设置提示: {e_embed}")
        # 17 represents wdFormatPDF
        doc.SaveAs(pdf_path, FileFormat=17)
        doc.Close()
        print(f"  ✓ Word COM 导出 PDF 成功，大小: {os.path.getsize(pdf_path):,} 字节")
    except Exception as e:
        print(f"  ! Word COM 导出失败: {e}")
        raise e
    finally:
        if word_app:
            try:
                word_app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    create_script_document()

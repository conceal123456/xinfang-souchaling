# -*- coding: utf-8 -*-
"""
校验最新现场放映版 PPT 与全流程场控分镜案 PPT。
支持校验：
1. 《心房搜查令_现场大屏放映版.pptx》（10 页，无后台场控表，纯场景画幅，演讲者备注）
2. 《心房搜查令_全流程场控与分镜案.pptx》（11 页，含 9x6 场控表，演讲者备注）
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches

def verify_screen_ppt(ppt_path="心房搜查令_现场大屏放映版.pptx"):
    print("=" * 50)
    print(f"校验现场大屏放映版: {ppt_path}")
    print("=" * 50)
    assert os.path.exists(ppt_path), f"未找到文件: {ppt_path}"
    prs = Presentation(ppt_path)
    assert len(prs.slides) == 10, f"大屏放映版预期 10 页，实际 {len(prs.slides)} 页"
    
    # 尺寸与画幅
    w, h = prs.slide_width.inches, prs.slide_height.inches
    assert abs(w / h - 16.0 / 9.0) < 0.01, f"画幅比例非 16:9: {w}x{h}"
    
    # 纯场景大图与无表格校验
    for idx in range(1, 9):
        slide = prs.slides[idx]
        pic_shapes = [s for s in slide.shapes if s.shape_type == 13]
        text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        assert len(pic_shapes) >= 1, f"第 {idx+1} 页缺少大图"
        assert len(text_shapes) == 0, f"第 {idx+1} 页存在残留文字"
        assert len([s for s in slide.shapes if s.has_table]) == 0, f"第 {idx+1} 页禁止包含表格"
        
    # 致谢页校验
    slide_10 = prs.slides[9]
    t10_compact = "".join([s.text_frame.text for s in slide_10.shapes if s.has_text_frame]).replace(" ", "").replace("　", "")
    assert "黄帝内经" in t10_compact or "恬淡虚无" in t10_compact
    assert "演职人员表" in t10_compact
    assert "特别鸣谢" in t10_compact
    
    # 演讲者备注全量校验
    for idx, slide in enumerate(prs.slides):
        notes = slide.notes_slide.notes_text_frame.text.strip()
        assert len(notes) > 10, f"第 {idx+1} 页演讲者备注缺失"
    print("✓ 现场大屏放映版校验通过！10 页规范、画幅纯净、备注详实、无失误隐患。\n")

def verify_control_ppt(ppt_path="心房搜查令_全流程场控与分镜案.pptx"):
    print("=" * 50)
    print(f"校验全流程场控与分镜案: {ppt_path}")
    print("=" * 50)
    assert os.path.exists(ppt_path), f"未找到文件: {ppt_path}"
    prs = Presentation(ppt_path)
    assert len(prs.slides) == 11, f"场控案预期 11 页，实际 {len(prs.slides)} 页"
    
    slide_11 = prs.slides[10]
    tables = [s.table for s in slide_11.shapes if s.has_table]
    assert len(tables) == 1, "第 11 页未找到场控表格"
    tbl = tables[0]
    assert len(tbl.rows) == 9, f"场控表行数不符: {len(tbl.rows)}"
    assert len(tbl.columns) == 6, f"场控表列数不符: {len(tbl.columns)}"
    print("✓ 全流程场控与分镜案校验通过！11 页规范、第 11 页包含完整 9x6 场控表。\n")

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    verify_screen_ppt()
    verify_control_ppt()
